"""Admin dashboard: products, tables, settings, uploads, stats, staff credentials, Excel export, backup/restore."""

import base64
import io
import json
import re
from datetime import datetime, time, timezone
from typing import Optional
from zoneinfo import ZoneInfo

from fastapi import APIRouter, Depends, HTTPException, Response
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from pymongo.errors import DuplicateKeyError

from lib.auth import hash_password, require_admin
from lib.db import db
from models.schemas import (BackupIn, BulkOut, CredentialsIn, DiningTable, Fulfillment, Message, MethodBreakdown,
                            Order, OrderStatus, PaymentMethod, Product, ProductIn, RestoreOut, Settings,
                            StaffAccount, Stats, TableBulkIn, TableIn, TableUpdate, TopItem, UploadIn, UploadOut,
                            new_id)
from routers.cashier import order_query, to_order
from routers.public import load_settings

router = APIRouter(prefix="/admin", dependencies=[Depends(require_admin)])
TZ = ZoneInfo("Asia/Manila")
MAX_IMAGE_BYTES = 4 * 1024 * 1024
ALLOWED_IMAGES = {"image/png", "image/jpeg", "image/webp", "image/gif"}
MAGIC = {"image/png": b"\x89PNG", "image/jpeg": b"\xff\xd8\xff", "image/gif": b"GIF8", "image/webp": b"RIFF"}
MAX_BULK_TABLES = 500


# ---------- Products ----------
@router.get("/products", response_model=list[Product])
async def all_products():
    docs = await db.products.find().sort([("sort_order", 1), ("name", 1)]).to_list(2000)
    return [Product(**d) for d in docs]


@router.post("/products", response_model=Product, status_code=201)
async def create_product(body: ProductIn):
    p = Product(**body.model_dump())
    await db.products.insert_one(p.model_dump())
    return p


@router.put("/products/{product_id}", response_model=Product)
async def update_product(product_id: str, body: ProductIn):
    doc = await db.products.find_one({"id": product_id})
    if not doc:
        raise HTTPException(status_code=404, detail="Product not found")
    await db.products.update_one({"id": product_id}, {"$set": body.model_dump()})
    return Product(**{**doc, **body.model_dump()})


@router.delete("/products/{product_id}", response_model=Message)
async def delete_product(product_id: str):
    res = await db.products.delete_one({"id": product_id})
    if not res.deleted_count:
        raise HTTPException(status_code=404, detail="Product not found")
    return Message(message="Product deleted")  # past orders keep their own copy of name/price


# ---------- Tables ----------
@router.get("/tables", response_model=list[DiningTable])
async def list_tables():
    docs = await db.tables.find().sort("number", 1).to_list(10000)
    return [DiningTable(**d) for d in docs]


@router.post("/tables", response_model=DiningTable, status_code=201)
async def add_table(body: TableIn):
    t = DiningTable(number=body.number, label=body.label.strip())
    try:
        await db.tables.insert_one(t.model_dump())
    except DuplicateKeyError:
        raise HTTPException(status_code=409, detail=f"Table {body.number} already exists")
    return t


@router.post("/tables/bulk", response_model=BulkOut, status_code=201)
async def bulk_add_tables(body: TableBulkIn):
    if body.end < body.start:
        raise HTTPException(status_code=400, detail="End number must be greater than or equal to start number")
    if body.end - body.start + 1 > MAX_BULK_TABLES:
        raise HTTPException(status_code=400, detail=f"You can add up to {MAX_BULK_TABLES} tables at a time")
    existing = {d["number"] for d in await db.tables.find({"number": {"$gte": body.start, "$lte": body.end}}).to_list(10000)}
    new = [DiningTable(number=n).model_dump() for n in range(body.start, body.end + 1) if n not in existing]
    if new:
        await db.tables.insert_many(new, ordered=False)
    return BulkOut(created=len(new), skipped=len(existing))


@router.put("/tables/{number}", response_model=DiningTable)
async def update_table(number: int, body: TableUpdate):
    doc = await db.tables.find_one_and_update({"number": number},
                                              {"$set": {"label": body.label.strip(), "active": body.active}},
                                              return_document=True)
    if not doc:
        raise HTTPException(status_code=404, detail="Table not found")
    return DiningTable(**doc)


# ---------- Settings ----------
@router.put("/settings", response_model=Settings)
async def save_settings(body: Settings):
    cats = [c.strip() for c in body.categories if c.strip()]
    body.categories = list(dict.fromkeys(cats)) or ["Coffee"]
    body.shop_name = body.shop_name.strip()
    await db.settings.update_one({"_key": "main"}, {"$set": body.model_dump()}, upsert=True)
    return body


# ---------- Uploads ----------
@router.post("/uploads", response_model=UploadOut, status_code=201)
async def upload_image(body: UploadIn):
    m = re.match(r"^data:(image/[a-zA-Z0-9.+-]+);base64,(.+)$", body.data_url, re.S)
    if not m or m.group(1) not in ALLOWED_IMAGES:
        raise HTTPException(status_code=400, detail="Only PNG, JPG, WEBP or GIF images are allowed")
    try:
        raw = base64.b64decode(m.group(2), validate=True)
    except Exception:
        raise HTTPException(status_code=400, detail="Corrupted image data")
    if len(raw) > MAX_IMAGE_BYTES:
        raise HTTPException(status_code=413, detail="Image must be 4 MB or smaller")
    if not raw.startswith(MAGIC[m.group(1)]):
        raise HTTPException(status_code=400, detail="File content does not match an image")
    image_id = new_id()
    await db.images.insert_one({"id": image_id, "content_type": m.group(1), "data": m.group(2),
                                "created_at": datetime.now(timezone.utc)})
    return UploadOut(url=f"/api/images/{image_id}")


# ---------- Staff accounts ----------
@router.get("/accounts", response_model=list[StaffAccount])
async def accounts():
    docs = await db.users.find().sort("role", 1).to_list(10)
    return [StaffAccount(role=d["role"], username=d["username"]) for d in docs]


@router.put("/accounts/{role}", response_model=StaffAccount)
async def update_account(role: str, body: CredentialsIn):
    if role not in ("admin", "cashier"):
        raise HTTPException(status_code=404, detail="Unknown role")
    clash = await db.users.find_one({"username": body.username, "role": {"$ne": role}})
    if clash:
        raise HTTPException(status_code=400, detail="Username already used by another account")
    await db.users.update_one({"role": role},
                              {"$set": {"username": body.username, "password_hash": hash_password(body.password)}},
                              upsert=True)
    if role == "cashier":
        await db.sessions.delete_many({"role": "cashier"})
    return StaffAccount(role=role, username=body.username)  # type: ignore[arg-type]


# ---------- Stats ----------
@router.get("/stats", response_model=Stats)
async def stats():
    now = datetime.now(TZ)
    start = datetime.combine(now.date(), time.min, TZ).astimezone(timezone.utc)
    paid = await db.orders.find({"status": "paid"}).to_list(100000)
    today_paid = [o for o in paid if o.get("paid_at") and o["paid_at"] >= start]
    items: dict[str, TopItem] = {}
    methods: dict[str, MethodBreakdown] = {}
    for o in paid:
        for it in o["items"]:
            t = items.setdefault(it["name"], TopItem(name=it["name"], qty=0, sales=0))
            t.qty += it["qty"]
            t.sales = round(t.sales + it["subtotal"], 2)
        m = methods.setdefault(o["payment_method"], MethodBreakdown(method=o["payment_method"], count=0, sales=0))
        m.count += 1
        m.sales = round(m.sales + o["total"], 2)
    return Stats(
        today=now.strftime("%Y-%m-%d"),
        today_sales=round(sum(o["total"] for o in today_paid), 2),
        today_orders=await db.orders.count_documents({"created_at": {"$gte": start}}),
        unpaid_orders=await db.orders.count_documents({"status": "unpaid"}),
        unserved_orders=await db.orders.count_documents({"status": {"$ne": "cancelled"}, "fulfillment": {"$ne": "served"}}),
        total_sales=round(sum(o["total"] for o in paid), 2),
        total_orders=len(paid),
        product_count=await db.products.count_documents({}),
        table_count=await db.tables.count_documents({"active": True}),
        top_items=sorted(items.values(), key=lambda t: t.qty, reverse=True)[:5],
        by_method=list(methods.values()),
    )


# ---------- Excel export ----------
PESO = '"PHP" #,##0.00'
HEAD_FILL = PatternFill("solid", fgColor="8C4A27")
HEAD_FONT = Font(bold=True, color="FFFFFF")
METHOD_NAME = {"cash": "Cash", "gcash": "GCash", "card": "Card / ATM"}


def local(dt: Optional[datetime]):
    return dt.astimezone(TZ).replace(tzinfo=None) if dt else None


def sheet(wb: Workbook, title: str, headers: list[str], widths: list[int], first: bool = False):
    ws = wb.active if first else wb.create_sheet(title)
    ws.title = title
    ws.append(headers)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font, c.alignment = HEAD_FILL, HEAD_FONT, Alignment(vertical="center")
    ws.freeze_panes = "A2"
    return ws


def fmt_cols(ws, money_cols: list[int], date_cols: list[int]):
    for row in ws.iter_rows(min_row=2):
        for c in row:
            if c.column in money_cols:
                c.number_format = PESO
            elif c.column in date_cols:
                c.number_format = "yyyy-mm-dd hh:mm"


@router.get("/export.xlsx")
async def export_excel(status: Optional[OrderStatus] = None, fulfillment: Optional[Fulfillment] = None,
                       method: Optional[PaymentMethod] = None, start: Optional[str] = None,
                       end: Optional[str] = None, include_products: bool = True):
    orders = [to_order(d) for d in
              await db.orders.find(order_query(status, fulfillment, method, None, start, end)).sort("created_at", 1).to_list(200000)]
    wb = Workbook()
    so = sheet(wb, "Order Summary",
               ["Order #", "Date and Time", "Table", "Table Label", "Customer", "Products", "Total Qty", "Order Total",
                "Payment Method", "Payment Status", "Order Status", "Notes", "Amount Tendered", "Change",
                "Cashier", "Paid At", "Served At"],
               [10, 18, 8, 14, 18, 50, 10, 14, 15, 15, 13, 30, 16, 12, 12, 18, 18], first=True)
    si = sheet(wb, "Order Items",
               ["Order #", "Date and Time", "Table", "Customer", "Product", "Quantity", "Unit Price", "Subtotal",
                "Payment Method", "Payment Status", "Order Status"],
               [10, 18, 8, 18, 28, 10, 14, 14, 15, 15, 13])
    daily: dict[str, list[float]] = {}
    for o in orders:
        so.append([o.order_number, local(o.created_at), o.table_number, o.table_label, o.nickname,
                   ", ".join(f"{i.qty} x {i.name}" for i in o.items), o.item_count, o.total,
                   METHOD_NAME[o.payment_method], o.status.capitalize(), o.fulfillment.capitalize(), o.notes,
                   o.amount_tendered, o.change, o.processed_by or "", local(o.paid_at), local(o.served_at)])
        for i in o.items:
            si.append([o.order_number, local(o.created_at), o.table_number, o.nickname, i.name, i.qty, i.price,
                       i.subtotal, METHOD_NAME[o.payment_method], o.status.capitalize(), o.fulfillment.capitalize()])
        if o.status == "paid":
            d = daily.setdefault(o.created_at.astimezone(TZ).strftime("%Y-%m-%d"), [0, 0, 0])
            d[0] += 1
            d[1] += o.item_count
            d[2] += o.total
    paid_total = round(sum(o.total for o in orders if o.status == "paid"), 2)
    so.append([])
    so.append(["", "", "", "", "", "TOTAL PAID", sum(o.item_count for o in orders if o.status == "paid"), paid_total])
    so.cell(row=so.max_row, column=6).font = Font(bold=True)
    fmt_cols(so, [8, 13, 14], [2, 16, 17])
    fmt_cols(si, [7, 8], [2])
    for row in so.iter_rows(min_row=2, min_col=6, max_col=6):
        row[0].alignment = Alignment(wrap_text=True, vertical="top")

    sd = sheet(wb, "Daily Sales", ["Date", "Paid Orders", "Items Sold", "Sales"], [14, 12, 12, 16])
    for day, (cnt, qty, total) in sorted(daily.items()):
        sd.append([day, cnt, qty, round(total, 2)])
    fmt_cols(sd, [4], [])

    if include_products:
        sp = sheet(wb, "Products", ["Name", "Category", "Price", "Available", "Description"], [28, 14, 12, 10, 50])
        for p in await db.products.find().sort("sort_order", 1).to_list(5000):
            sp.append([p["name"], p["category"], p["price"], "Yes" if p["available"] else "No", p.get("description", "")])
        fmt_cols(sp, [3], [])

    buf = io.BytesIO()
    wb.save(buf)
    rng = f"{start or 'all'}_to_{end or 'now'}"
    name = f"coffee-beats-orders-{rng}.xlsx"
    return Response(content=buf.getvalue(),
                    media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": f'attachment; filename="{name}"', "X-Order-Count": str(len(orders)),
                             "Access-Control-Expose-Headers": "X-Order-Count, Content-Disposition"})


# ---------- Backup / restore ----------
def clean(doc: dict) -> dict:
    doc.pop("_id", None)
    doc.pop("_key", None)
    return doc


@router.get("/backup")
async def backup():
    payload = {
        "version": 2,
        "exported_at": datetime.now(timezone.utc).isoformat(),
        "settings": (await load_settings()).model_dump(),
        "products": [clean(d) for d in await db.products.find().to_list(10000)],
        "tables": [clean(d) for d in await db.tables.find().to_list(10000)],
        "orders": [clean(d) for d in await db.orders.find().to_list(200000)],
        "images": [clean(d) for d in await db.images.find().to_list(10000)],
    }
    name = f"coffee-beats-backup-{datetime.now(TZ).strftime('%Y%m%d-%H%M')}.json"
    return Response(content=json.dumps(payload, default=str, indent=1), media_type="application/json",
                    headers={"Content-Disposition": f'attachment; filename="{name}"'})


def parse_dt(v):
    if isinstance(v, str) and v:
        try:
            dt = datetime.fromisoformat(v)
            return dt if dt.tzinfo else dt.replace(tzinfo=timezone.utc)
        except ValueError:
            return None
    return v


@router.post("/restore", response_model=RestoreOut)
async def restore(body: BackupIn):
    try:
        products = [Product(**p).model_dump() for p in body.products]
        orders = []
        for o in body.orders:
            m = Order(**o).model_dump()
            if o.get("client_token"):
                m["client_token"] = o["client_token"]
            orders.append(m)
        tables = [DiningTable(**t).model_dump() for t in body.tables]
        settings = Settings(**{k: v for k, v in body.settings.items() if k in Settings.model_fields}).model_dump() \
            if body.settings else None
    except Exception as exc:
        raise HTTPException(status_code=400, detail=f"Invalid backup file: {str(exc)[:200]}")
    if len({t["number"] for t in tables}) != len(tables):
        raise HTTPException(status_code=400, detail="Backup contains duplicate table numbers")
    images = []
    for img in body.images:
        if {"id", "content_type", "data"} <= img.keys() and img["content_type"] in ALLOWED_IMAGES:
            images.append({"id": img["id"], "content_type": img["content_type"], "data": img["data"],
                           "created_at": parse_dt(img.get("created_at")) or datetime.now(timezone.utc)})
    await db.products.delete_many({})
    await db.orders.delete_many({})
    if products:
        await db.products.insert_many(products)
    if orders:
        await db.orders.insert_many(orders)
        top = max(o["order_number"] for o in orders)
        await db.counters.update_one({"_id": "order_number"}, {"$max": {"seq": top}}, upsert=True)
    if tables:
        await db.tables.delete_many({})
        await db.tables.insert_many(tables)
    for img in images:
        await db.images.update_one({"id": img["id"]}, {"$set": img}, upsert=True)
    if settings:
        await db.settings.update_one({"_key": "main"}, {"$set": settings}, upsert=True)
    return RestoreOut(products=len(products), orders=len(orders), images=len(images), tables=len(tables))


